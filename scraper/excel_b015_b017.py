#!/usr/bin/env python3
"""Excel PERSONALE solo classi B015 e B017. Scritto FUORI dal repository (non va mai pubblicato)."""
import json, sys
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else Path.home() / "personale" / "Interpellati_B015_B017_PERSONALE.xlsx"
COL = {"verde": "63BE7B", "arancione": "FFC000", "rosso": "F8696B"}
J = lambda p, d: json.loads(p.read_text("utf-8")) if p.exists() else d

mem = J(ROOT / "data/memoria_scuole.json", {}).get("scuole", {})
sede2ist = {}
for ep in [ROOT / "data/sicilia_scuole.json"] + sorted((ROOT / "data/scuole").glob("*_elenco.json")):
    if "priorita" in ep.name:
        continue
    for cod, v in J(ep, {}).get("scuole", {}).items():
        sede2ist[cod] = cod
        for s in v.get("sedi", []):
            sede2ist[s] = cod
feed = J(ROOT / "data/feed.json", {}).get("interpelli", [])
arch = J(ROOT / "data/interpelli.json", {}).get("interpelli", [])
def bs(it): return any(c in ("B015", "B017", "B-15", "B-17") for c in (it.get("classi") or []))
bint = [i for i in {x.get("id") or x.get("link"): x for x in arch + feed}.values() if bs(i)]
per_ist = {}
for i in bint:
    per_ist.setdefault(i.get("codiceScuola"), []).append(i)

src = load_workbook(ROOT / "data/Elenco_scuole_B015_B017.xlsx", read_only=True)["Tutta Italia"]
rows = list(src.iter_rows(values_only=True)); head = rows[0]; rows = rows[1:]
wb = Workbook(); ws = wb.active; ws.title = "Scuole B015-B017"
cols = ["Regione", "Provincia", "Comune", "Scuola", "Codice", "Classi", "Verifica", "Tipologia", "Indirizzo", "Email", "PEC", "Sito",
        "Istituto di riferimento", "Stato lettura sito", "Motivo", "Interpelli B015/B017 trovati", "Ultimo interpello", "Link ultimo"]
ws.append(cols)
hf, hfill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3864")
for c in ws[1]: c.font, c.fill, c.alignment = hf, hfill, Alignment(wrap_text=True, vertical="center")
tot = {"verde": 0, "arancione": 0, "rosso": 0, "n.d.": 0}
H = {h: i for i, h in enumerate(head)}
for r in sorted(rows, key=lambda r: (r[0] or "", r[1] or "", r[2] or "", r[3] or "")):
    ist = sede2ist.get(r[H["Codice"]])
    m = mem.get(ist) if ist else None
    st = m["stato"] if m else "n.d."
    tot[st] += 1
    its = sorted(per_ist.get(ist, []), key=lambda i: i.get("pubblicato") or "")
    u = its[-1] if its else {}
    ws.append([r[H["Regione"]], r[H["Provincia"]], r[H["Comune"]], r[H["Scuola"]], r[H["Codice"]], r[H["Classi"]], r[H["Verifica"]],
               r[H["Tipologia"]], r[H["Indirizzo"]], r[H["Email"]], r[H["PEC"]], (m or {}).get("sito") or r[H["Sito"]], r[H["Istituto di riferimento"]],
               st.upper(), (m or {}).get("motivo", "non in anagrafica ministeriale"), len(its), u.get("pubblicato", ""), u.get("link", "")])
    if st in COL:
        ws.cell(ws.max_row, 14).fill = PatternFill("solid", fgColor=COL[st]); ws.cell(ws.max_row, 14).font = Font(bold=True)
for i, w in enumerate([16, 9, 20, 38, 13, 11, 11, 24, 32, 28, 28, 36, 32, 14, 48, 12, 14, 40], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions

w2 = wb.create_sheet("Interpelli B015-B017")
w2.append(["Scuola", "Codice", "Comune", "Classi", "Titolo", "Pubblicato", "Scadenza", "Riservato ai docenti interni", "Link"])
for c in w2[1]: c.font, c.fill = hf, hfill
for i in sorted(bint, key=lambda i: i.get("pubblicato") or "", reverse=True):
    w2.append([i.get("scuola", ""), i.get("codiceScuola", ""), i.get("comune", ""), ", ".join(i.get("classi") or []), i.get("titolo", ""),
               i.get("pubblicato", ""), i.get("scadenza", ""), "Sì" if i.get("riservato") == "interni" else "", i.get("link", "")])
for i, w in enumerate([36, 13, 18, 14, 60, 12, 20, 12, 50], 1): w2.column_dimensions[get_column_letter(i)].width = w
w2.freeze_panes = "A2"

w0 = wb.create_sheet("Riepilogo", 0)
w0["A1"] = "Classi B015 / B017 – scuole da controllare (file personale)"; w0["A1"].font = Font(bold=True, size=14)
w0["A2"] = f"Aggiornato il {datetime.now():%d/%m/%Y %H:%M} · verde = sito letto · arancione = ancora in ricerca · rosso = non si trova"
w0.append([]); w0.append(["Stato", "Scuole"])
for c in w0[4]: c.font, c.fill = hf, hfill
for k, et in (("verde", "Verde – sito letto"), ("arancione", "Arancione – in ricerca"), ("rosso", "Rosso – non si trova"), ("n.d.", "Non in anagrafica")):
    w0.append([et, tot[k]])
    if k in COL: w0.cell(w0.max_row, 1).fill = PatternFill("solid", fgColor=COL[k])
w0.append(["Totale", len(rows)]); w0.append(["Interpelli B015/B017 trovati", len(bint)])
w0.column_dimensions["A"].width = 40
OUT.parent.mkdir(parents=True, exist_ok=True); wb.save(OUT)
print(OUT, tot, len(bint))
