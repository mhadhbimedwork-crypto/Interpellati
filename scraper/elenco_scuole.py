"""Elenco Excel delle scuole italiane dove possono servire le classi B015 / B017.

Fonte: open data MIM (anagrafica scuole statali + indirizzi di studio). Solo dati pubblici delle scuole
(denominazione, indirizzo, email/PEC istituzionali, sito). Ordinato Regione -> Provincia -> Comune -> Scuola.
"""
from __future__ import annotations

import csv
import io
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import geo

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "data" / "Elenco_scuole_B015_B017.xlsx"
REP = ROOT / "data" / "elenco_scuole_report.txt"
LOG: list[str] = []


def log(*a):
    s = " ".join(str(x) for x in a)
    LOG.append(s)
    print(s, flush=True)


def scarica(nome: str, anni: list[str]):
    s = requests.Session()
    s.headers["User-Agent"] = geo.UA
    for anno in anni:
        url = f"{geo.BASE}{nome}{anno}.csv"
        try:
            r = s.get(url, timeout=300)
        except requests.RequestException as e:
            log(f"  {nome}{anno}: {type(e).__name__}")
            continue
        log(f"  {nome}{anno}: HTTP {r.status_code} {len(r.content)} byte")
        if r.status_code == 200 and len(r.content) > 5000:
            for enc in ("utf-8-sig", "latin-1"):
                try:
                    return anno, list(csv.DictReader(io.StringIO(r.content.decode(enc))))
                except UnicodeDecodeError:
                    pass
    return None, []


def col(row: dict, *chiavi: str) -> str:
    for k in chiavi:
        v = row.get(k)
        if v and v.strip() and v.strip().upper() not in ("NON DISPONIBILE", "ND"):
            return v.strip()
    return ""


B015 = re.compile(r"elettronic|elettrotecnic|elettric|automazion|telecomunicazion|manutenzione e assistenza|industria e artigianato|energia|domotic|impianti", re.I)
B017 = re.compile(r"meccanic|meccatronic|energia|manutenzione e assistenza|industria e artigianato|produzioni industriali|trasporti e logistica|costruzione del mezzo|conduzione del mezzo|mezzi di trasporto", re.I)
TECPRO = re.compile(r"tecnic|professional", re.I)
NOME_TEC = re.compile(r"\b(ITIS|ITI|IPSIA|IPIA|ITT|ITE|IIS|ISIS|IS|I\.?T\.?I\.?S?\.?|I\.?P\.?S\.?I\.?A\.?|TECNIC|PROFESSIONAL|INDUSTRIAL|ARTIGIANAT)\b", re.I)


def sito(v: str) -> str:
    v = v.strip()
    if not v:
        return ""
    v = re.sub(r"^(https?)[:;]?/{1,3}", r"\1://", v, flags=re.I)
    return v if v.lower().startswith("http") else "https://" + v


def main() -> int:
    log("Anagrafica scuole (statali + province autonome)")
    righe: list[dict] = []
    for nome in ("SCUANAGRAFESTAT", "SCUANAAUTSTAT"):
        anno, rows = scarica(nome, geo.ANNI)
        if rows:
            log("  colonne:", ",".join(rows[0].keys()))
        righe += rows
    if not righe:
        log("ERRORE: anagrafica non scaricata")
        REP.write_text("\n".join(LOG) + "\n", "utf-8")
        return 1

    # indirizzi di studio per scuola (dataset alunni/classi per indirizzo)
    ind: dict[str, set[str]] = defaultdict(set)
    for nome in ("ALUCORSOINDCLASTA",):
        anno, rows = scarica(nome, ["20242520250831", "20232420240831", "20252620260831"])
        if rows:
            log("  colonne indirizzi:", ",".join(rows[0].keys()))
            for r in rows:
                cod = (r.get("CODICESCUOLA") or "").strip().upper()
                txt = " | ".join(v for k, v in r.items() if v and re.search(r"INDIRIZZ|ARTICOLAZ|DESCRIZ", k, re.I) and not re.search(r"INDIRIZZOSCUOLA|EMAIL|PEC", k, re.I))
                if cod and txt:
                    ind[cod].add(txt)
            log(f"  scuole con indirizzi noti: {len(ind)}")
            break

    # interpelli B015/B017 gia' visti dallo scraper
    avvisi: dict[str, list[dict]] = defaultdict(list)
    try:
        for i in json.loads((ROOT / "data" / "interpelli.json").read_text("utf-8"))["interpelli"]:
            if i.get("codiceScuola") and set(i.get("classi", [])) & {"B015", "B017"}:
                avvisi[i["codiceScuola"].upper()].append(i)
    except (OSError, ValueError, KeyError):
        pass

    # una riga per scuola con indirizzi tecnico-professionali (secondaria II grado)
    out = []
    for r in righe:
        cod = col(r, "CODICESCUOLA").upper()
        tip = col(r, "DESCRIZIONETIPOLOGIAGRADOISTRUZIONESCUOLA")
        nome = col(r, "DENOMINAZIONESCUOLA")
        rif = col(r, "CODICEISTITUTORIFERIMENTO").upper() or cod
        testo_ind = " ; ".join(sorted(ind.get(cod, set()) | (ind.get(rif, set()) if rif != cod else set())))
        cl = []
        if B015.search(testo_ind):
            cl.append("B015")
        if B017.search(testo_ind):
            cl.append("B017")
        da_verificare = False
        if not cl:
            sec2 = re.search(r"II\s*GRADO|SECONDARIA\s+DI\s+II|TECNIC|PROFESSIONAL|SUPERIORE", tip, re.I)
            if (sec2 and TECPRO.search(tip + " " + nome)) or (sec2 and NOME_TEC.search(nome)):
                cl, da_verificare = ["B015?", "B017?"], True
        if not cl:
            continue
        av = avvisi.get(cod, []) + avvisi.get(rif, [])
        out.append({
            "Regione": col(r, "REGIONE").title(), "Provincia": col(r, "PROVINCIA").title(),
            "Comune": col(r, "DESCRIZIONECOMUNE").title(), "Scuola": nome, "Codice": cod,
            "Tipologia": tip.title(), "Classi": "/".join(cl), "Verifica": "da verificare" if da_verificare else "indirizzi confermati",
            "Indirizzi di studio": testo_ind[:600], "Indirizzo": col(r, "INDIRIZZOSCUOLA"), "CAP": col(r, "CAPSCUOLA"),
            "Email": col(r, "INDIRIZZOEMAILSCUOLA").lower(), "PEC": col(r, "INDIRIZZOPECSCUOLA").lower(),
            "Sito": sito(col(r, "SITOWEBSCUOLA")), "Istituto di riferimento": col(r, "DENOMINAZIONEISTITUTORIFERIMENTO"),
            "Interpelli B015/B017 visti": len(av),
            "Ultimo interpello": max((a.get("pubblicato", "") for a in av), default=""),
            "Link ultimo": next((a["link"] for a in sorted(av, key=lambda a: a.get("pubblicato", ""), reverse=True)), ""),
        })
    out.sort(key=lambda x: (x["Regione"], x["Provincia"], x["Comune"], x["Scuola"]))
    log(f"Scuole in elenco: {len(out)} (confermate da indirizzi: {sum(1 for x in out if x['Verifica'].startswith('ind'))})")
    log("Per regione:", dict(Counter(x["Regione"] for x in out)))

    cols = ["Regione", "Provincia", "Comune", "Scuola", "Codice", "Classi", "Verifica", "Tipologia", "Indirizzi di studio",
            "Indirizzo", "CAP", "Email", "PEC", "Sito", "Istituto di riferimento",
            "Interpelli B015/B017 visti", "Ultimo interpello", "Link ultimo"]
    larg = [16, 14, 18, 38, 12, 11, 18, 22, 50, 30, 7, 30, 30, 34, 30, 12, 12, 40]
    testa = PatternFill("solid", fgColor="1F4E78")

    def foglio(ws, rows):
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = testa
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for x in rows:
            ws.append([x[c] for c in cols])
        for i, w in enumerate(larg, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "E2"
        ws.auto_filter.ref = ws.dimensions
        for row in ws.iter_rows(min_row=2):
            for cell, nome in zip(row, cols):
                if nome == "Sito" and cell.value:
                    cell.hyperlink = cell.value
                    cell.font = Font(color="0563C1", underline="single")
                elif nome == "Link ultimo" and cell.value:
                    cell.hyperlink = cell.value
                    cell.font = Font(color="0563C1", underline="single")
                elif nome == "Email" and cell.value:
                    cell.hyperlink = "mailto:" + cell.value
                    cell.font = Font(color="0563C1", underline="single")

    wb = Workbook()
    ws = wb.active
    ws.title = "Riepilogo"
    ws.append(["Elenco scuole per classi di concorso B015 / B017"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Fonte: open data MIM (anagrafica scuole statali + indirizzi di studio). 'indirizzi confermati' = la scuola ha un indirizzo elettrico/elettronico/meccanico; 'da verificare' = istituto tecnico/professionale senza dato sugli indirizzi."])
    ws.append(["Le classi indicano dove i laboratori sono plausibili: la reale disponibilita' di cattedre/ore si vede solo negli interpelli e nelle MAD."])
    ws.append([])
    ws.append(["Regione", "Scuole", "di cui confermate", "Province", "Comuni", "Con interpelli B015/B017 visti"])
    for c in ws[5]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = testa
    per = defaultdict(list)
    for x in out:
        per[x["Regione"]].append(x)
    for reg in sorted(per):
        v = per[reg]
        ws.append([reg, len(v), sum(1 for x in v if x["Verifica"].startswith("ind")), len({x["Provincia"] for x in v}),
                   len({x["Comune"] for x in v}), sum(1 for x in v if x["Interpelli B015/B017 visti"])])
    ws.append(["TOTALE", len(out), sum(1 for x in out if x["Verifica"].startswith("ind")), "", "", sum(1 for x in out if x["Interpelli B015/B017 visti"])])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for i, w in enumerate([22, 10, 18, 10, 10, 28], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    foglio(wb.create_sheet("Tutta Italia"), out)
    if "Sicilia" in per:
        foglio(wb.create_sheet("SICILIA"), per["Sicilia"])
    for reg in sorted(per):
        if reg != "Sicilia":
            foglio(wb.create_sheet(reg[:31]), per[reg])
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    REP.write_text("\n".join(LOG) + "\n", "utf-8")
    log("Salvato", OUT, OUT.stat().st_size, "byte")
    return 0


if __name__ == "__main__":
    sys.exit(main())
