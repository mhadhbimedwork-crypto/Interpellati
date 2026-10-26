"""Excel principale: tutte le scuole e tutte le classi di concorso, ordinati per regione > provincia > comune.
Stato di ogni scuola: TROVATA (verde) · DA CERCARE (arancione) · INESISTENTE (rosso). Solo dati pubblici."""
from __future__ import annotations
import json
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
J = lambda p, d: json.loads(p.read_text("utf-8")) if p.exists() else d
COL = {"verde": "63BE7B", "arancione": "FFC000", "rosso": "F8696B"}
CHI = {"verde": "C6EFCE", "arancione": "FFE699", "rosso": "FFC7CE"}
ET = {"verde": "TROVATA", "arancione": "DA CERCARE", "rosso": "INESISTENTE"}
HF, HFILL = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3864")


def azione(e):
    m = e["motivo"]
    if e["stato"] == "verde": return "Leggere ogni giorno (già a posto)"
    if "nessun sito" in m: return "Cercare il sito su Indice PA / ufficio scolastico provinciale"
    if "dominio inesistente" in m: return "Cercare il nuovo dominio (Indice PA, ricerca per nome e comune)"
    if "non ancora pubblicato" in m: return "Ricontrollare ogni giorno finché il sito è pubblicato"
    if "robots" in m: return "Non leggere il sito (robots.txt): cercare avvisi su UST/USR e aggregatori"
    if "limita" in m: return "Rileggere a ritmo lento (server condiviso)"
    return "Riprovare la lettura del sito oggi"


def foglio(wb, nome, intest, righe, larg, stato_col=None, wrap=()):
    ws = wb.create_sheet(nome); ws.append(intest)
    for c in ws[1]: c.font, c.fill, c.alignment = HF, HFILL, Alignment(wrap_text=True, vertical="center")
    for r in righe:
        ws.append(r)
        if stato_col:
            k = {v: k for k, v in ET.items()}.get(str(r[stato_col - 1]))
            if k: ws.cell(ws.max_row, stato_col).fill = PatternFill("solid", fgColor=COL[k]); ws.cell(ws.max_row, stato_col).font = Font(bold=True)
    for i, w in enumerate(larg, 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    return ws


def scrivi(mem, prog, out: Path):
    m = mem["scuole"]
    arch = {x.get("id") or x.get("link"): x for x in J(ROOT / "data/interpelli.json", {}).get("interpelli", []) + J(ROOT / "data/feed.json", {}).get("interpelli", [])}
    its = list(arch.values())
    cat = J(ROOT / "data/classi_concorso.json", {}).get("voci", [])
    alias = {}
    for v in cat: alias[v["codice"]] = v["codice"]; alias[v["alias"]] = v["codice"]
    per_sc = defaultdict(list)
    for i in its:
        if i.get("codiceScuola"): per_sc[i["codiceScuola"]].append(i)
    tot = Counter(e["stato"] for e in m.values())
    n = len(m)
    wb = Workbook(); w0 = wb.active; w0.title = "Riepilogo"
    w0["A1"] = "Interpellati – scuole e classi di concorso d'Italia"; w0["A1"].font = Font(bold=True, size=14)
    w0["A2"] = f"Aggiornato il {datetime.now():%d/%m/%Y %H:%M} · TROVATA = sito letto · DA CERCARE = ancora in ricerca · INESISTENTE = non si trova (dopo 7 giorni di tentativi o sito assente)"
    w0.append([]); w0.append(["Stato", "Scuole", "%"])
    for c in w0[4]: c.font, c.fill = HF, HFILL
    for k in ("verde", "arancione", "rosso"):
        w0.append([ET[k], tot[k], round(100 * tot[k] / max(1, n), 1)]); w0.cell(w0.max_row, 1).fill = PatternFill("solid", fgColor=COL[k])
    w0.append(["Totale scuole nell'anagrafica", n, 100]); w0.append([])
    w0.append(["Regione", "Trovate", "Da cercare", "Inesistenti", "Totale", "% trovate"])
    for c in w0[w0.max_row]: c.font, c.fill = HF, HFILL
    reg = defaultdict(Counter)
    for e in m.values(): reg[e["regione"]][e["stato"]] += 1
    for r, c in sorted(reg.items()):
        t = sum(c.values()); w0.append([r, c["verde"], c["arancione"], c["rosso"], t, round(100 * c["verde"] / t, 1)])
        for j, k in ((2, "verde"), (3, "arancione"), (4, "rosso")): w0.cell(w0.max_row, j).fill = PatternFill("solid", fgColor=CHI[k])
    w0.append([]); w0.append(["Perché non sono trovate", "Scuole"])
    for c in w0[w0.max_row]: c.font, c.fill = HF, HFILL
    for k, v in Counter(e["motivo"].split(" (")[0] for e in m.values() if e["stato"] != "verde").most_common(): w0.append([k, v])
    w0.append([]); w0.append(["Storico", "Trovate", "Da cercare", "Inesistenti", "% trovate"])
    for c in w0[w0.max_row]: c.font, c.fill = HF, HFILL
    for r in prog["storico"][-60:][::-1]: w0.append([r["data"], r["verdi"], r["arancioni"], r["rossi"], r["percentuale_verdi"]])
    w0.column_dimensions["A"].width = 62
    for col in "BCDEF": w0.column_dimensions[col].width = 13

    ordine = sorted(m.items(), key=lambda x: (x[1]["regione"], x[1]["provincia"], x[1]["comune"], x[1]["nome"]))
    righe = []
    for cod, e in ordine:
        lst = sorted(per_sc.get(cod, []), key=lambda i: i.get("pubblicato") or "")
        classi = sorted({alias.get(c, c) for i in lst for c in i.get("classi", [])})
        u = lst[-1] if lst else {}
        righe.append([e["regione"], e["provincia"], e["comune"], e["nome"], cod, ET[e["stato"]], e["motivo"], azione(e), e.get("sito") or "",
                      e.get("ultima_lettura", ""), len(e.get("giorni_errore", [])), ", ".join(classi), len(lst), u.get("pubblicato", ""), u.get("link", ""), e.get("primo_visto", "")])
    foglio(wb, "Scuole", ["Regione", "Provincia", "Comune", "Scuola", "Codice", "Stato", "Motivo", "Cosa fare / dove leggere", "Sito", "Ultima lettura riuscita", "Giorni in errore",
                          "Classi viste negli avvisi", "Avvisi trovati", "Ultimo avviso", "Link ultimo avviso", "In elenco dal"],
           righe, [16, 9, 20, 42, 13, 14, 50, 50, 38, 14, 10, 30, 9, 12, 40, 12], stato_col=6)
    for st, nome in (("DA CERCARE", "Da cercare (oggi)"), ("INESISTENTE", "Inesistenti")):
        foglio(wb, nome, ["Regione", "Provincia", "Comune", "Scuola", "Codice", "Stato", "Motivo", "Cosa fare / dove leggere", "Sito", "Giorni in errore"],
               [[r[0], r[1], r[2], r[3], r[4], r[5], r[6], r[7], r[8], r[10]] for r in righe if r[5] == st], [16, 9, 20, 42, 13, 14, 50, 50, 38, 10], stato_col=6)

    # classi di concorso: tutte quelle d'Italia, con cosa e' stato trovato
    ci = defaultdict(list)
    for i in its:
        for c in i.get("classi", []): ci[alias.get(c, c)].append(i)
    rc = []
    for v in cat:
        l = ci.get(v["codice"], [])
        rc.append([v["gruppo"], v["codice"], v["alias"], v["nome"], len(l), len({i.get("regione") for i in l if i.get("regione")}), len({i.get("provincia") for i in l if i.get("provincia")}),
                   ", ".join(sorted({i.get("regione") for i in l if i.get("regione")})), max([i.get("pubblicato") or "" for i in l], default="")])
    foglio(wb, "Classi di concorso", ["Gruppo", "Codice", "Alias", "Denominazione", "Avvisi trovati", "Regioni", "Province", "Dove trovata", "Ultimo avviso"], rc, [34, 10, 9, 70, 10, 8, 8, 60, 12])

    # classe x regione e classe x provincia x comune (solo dove c'e' un avviso)
    reg_list = sorted({m[c]["regione"] for c in m})
    mat = []
    for v in cat:
        l = ci.get(v["codice"], [])
        if not l: continue
        cr = Counter(i.get("regione") for i in l)
        mat.append([v["codice"], v["nome"]] + [cr.get(r, 0) or "" for r in reg_list] + [len(l)])
    foglio(wb, "Classi per regione", ["Classe", "Denominazione"] + reg_list + ["Totale"], mat, [10, 50] + [9] * len(reg_list) + [9])
    cpc = Counter()
    for i in its:
        for c in i.get("classi", []): cpc[(alias.get(c, c), i.get("regione") or "", i.get("provincia") or "", i.get("comune") or "")] += 1
    nome_cl = {v["codice"]: v["nome"] for v in cat}
    foglio(wb, "Classi per comune", ["Classe", "Denominazione", "Regione", "Provincia", "Comune", "Avvisi"],
           [[c, nome_cl.get(c, ""), r, p, co, n_] for (c, r, p, co), n_ in sorted(cpc.items(), key=lambda x: (x[0][0], x[0][1], x[0][2], x[0][3]))], [10, 50, 16, 9, 22, 8])

    stato_sc = {c: ET[e["stato"]] for c, e in m.items()}
    foglio(wb, "Avvisi", ["Pubblicato", "Scadenza", "Classi", "Regione", "Provincia", "Comune", "Scuola", "Codice", "Stato scuola", "Riservato interni", "Fonte", "Link"],
           [[i.get("pubblicato", ""), i.get("scadenza", ""), ", ".join(i.get("classi", [])), i.get("regione", ""), i.get("provincia", ""), i.get("comune", ""), i.get("scuola", "")[:60],
             i.get("codiceScuola", ""), stato_sc.get(i.get("codiceScuola"), ""), "Sì" if i.get("riservato") == "interni" else "", i.get("fonte", ""), i.get("link", "")]
            for i in sorted(its, key=lambda i: i.get("pubblicato") or "", reverse=True)], [12, 12, 18, 16, 9, 20, 40, 13, 14, 9, 16, 50], stato_col=9)
    out.parent.mkdir(parents=True, exist_ok=True); wb.save(out)
