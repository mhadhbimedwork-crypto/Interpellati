#!/usr/bin/env python3
"""Scopre dall'Indice PA (fonte ufficiale AgID) i siti di tutti gli uffici scolastici: USR, UST/ATP/ambiti territoriali.
Scrive data/uffici_ipa.json = [{nome, url, provincia, regione}]. Uso: python scraper/uffici_ipa.py"""
from __future__ import annotations
import io, json, re, sys
from pathlib import Path
import requests
sys.path.insert(0, str(Path(__file__).resolve().parent))
import scrape as S  # noqa: E402
from diagnosi_siti import IPA_XLSX  # noqa: E402

OUT = S.ROOT / "data" / "uffici_ipa.json"
RX = re.compile(r"ambito territoriale|ufficio scolastico|ufficio\s+\w*\s*ambito|ufficio.*provinciale|\bATP\b|\bUST\b|\bUSP\b|\bUSR\b|centro servizi amministrativi|\bCSA\b|provveditorato", re.I)
NO = re.compile(r"istituto|liceo|scuola|convitto|circolo didattico|direzione didattica|conservatorio|accademia|universit", re.I)


def main() -> int:
    from openpyxl import load_workbook
    r = requests.get(IPA_XLSX, timeout=180, headers={"User-Agent": S.UA})
    r.raise_for_status()
    ws = load_workbook(io.BytesIO(r.content), read_only=True).active
    righe = ws.iter_rows(values_only=True)
    testa = [str(x or "") for x in next(righe)]
    print("colonne:", testa)
    i_n = next(i for i, h in enumerate(testa) if h.lower().startswith("denominazione_ente"))
    i_s = testa.index("Sito_istituzionale")
    i_c = next((i for i, h in enumerate(testa) if h.lower() in ("comune", "des_comune")), None)
    out, seen = [], set()
    for row in righe:
        nome, sito = str(row[i_n] or ""), str(row[i_s] or "").strip()
        if not sito or not RX.search(nome) or NO.search(nome):
            continue
        url = sito if "://" in sito else "https://" + sito
        if url in seen:
            continue
        seen.add(url)
        prov = S.find_province(nome) or (S.find_province(str(row[i_c] or "")) if i_c is not None else None)
        it = {"nome": nome, "url": url, "provincia": prov}
        if prov:
            it["regione"] = S.PROVINCE[prov][1]
        out.append(it)
    OUT.write_text(json.dumps(out, ensure_ascii=False, indent=0) + "\n", "utf-8")
    print(f"uffici scoperti: {len(out)} · con provincia: {sum(1 for o in out if o['provincia'])}")
    for o in out[:200]:
        print(" ", o["provincia"], "|", o["nome"][:70], "|", o["url"])
    return 0


if __name__ == "__main__":
    sys.exit(main())
