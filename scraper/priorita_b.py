#!/usr/bin/env python3
"""Ricerca potenziata per le classi B015 e B017: legge, a ogni giro, TUTTI gli istituti dove queste classi sono
plausibili (tecnici industriali, istituti superiori con sede tecnica, professionali) invece di una rotazione lenta.
Usa lo stesso motore delle scuole (home + sezioni interpelli/supplenze/MAD/albo/avvisi), bot dichiarato, robots rispettato.
Scrive data/scuole/prioritaB_*.json; la scansione principale importa gli avvisi trovati."""
from __future__ import annotations
import json, os, sys, time
from datetime import date, datetime, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import scrape as S  # noqa: E402
import sicilia  # noqa: E402
import scuole_italia as SI  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
D = ROOT / "data" / "scuole"


def codici_b() -> set[str]:
    wb = load_workbook(ROOT / "data" / "Elenco_scuole_B015_B017.xlsx", read_only=True)
    ws = wb["Tutta Italia"]
    out = set()
    for r in list(ws.iter_rows(values_only=True))[1:]:
        if r and r[4]:
            out.add(str(r[4]).strip().upper())
    return out


def main() -> int:
    b = codici_b()
    scelti: dict[str, dict] = {}
    for area, (nome, sigle) in SI.AREE.items():
        p = D / f"{area}_elenco.json"
        if not p.exists():
            continue
        sicilia.configura(area, tuple(sigle.split()), nome)
        sc = sicilia.carica_elenco(S.log)  # con i siti corretti
        for cod, v in sc.items():
            if cod in b or any(s in b for s in v.get("sedi", [])):
                scelti[cod] = v
    # Sicilia ha i suoi file storici
    ps = ROOT / "data" / "sicilia_scuole.json"
    if ps.exists():
        sicilia.ELENCO = ps
        sicilia.CORRETTI = ROOT / "data" / "sicilia_siti_corretti.json"
        sicilia.PREFISSI = sicilia.PREFISSI
        try:
            sc = sicilia._applica_correzioni(json.loads(ps.read_text("utf-8")).get("scuole", {}))
            for cod, v in sc.items():
                if cod in b or any(s in b for s in v.get("sedi", [])):
                    scelti[cod] = v
        except ValueError:
            pass
    S.log(f"B015/B017: istituti da leggere a ogni giro {len(scelti)} (con sito {sum(1 for v in scelti.values() if v.get('sito'))})")
    # motore su file propri
    sicilia.ELENCO, sicilia.STATO, sicilia.REPORT, sicilia.CORRETTI = (
        D / "prioritaB_elenco.json", D / "prioritaB_stato.json", D / "prioritaB_report.json", D / "prioritaB_corretti.json")
    sicilia.NOME_AREA = "Priorita B015/B017"
    (D / "prioritaB_elenco.json").write_text(json.dumps({"ts": time.time(), "scuole": scelti}, ensure_ascii=False, separators=(",", ":")) + "\n", "utf-8")
    sicilia.MAX_PER_RUN = 5000
    sicilia.THREADS = 32
    sicilia.MAX_SOTTO = 8
    try:
        store = json.loads((ROOT / "data" / "interpelli.json").read_text("utf-8"))
        sicilia.KNOWN = {i.get("link") for i in store.get("interpelli", [])}
    except (OSError, ValueError):
        sicilia.KNOWN = set()
    f = S.Fetcher()
    since = date.today() - timedelta(days=10)
    trovati = sicilia.scansiona_sicilia(S, f, since, budget_s=int(os.environ.get("BUDGET_S", "1200")))
    tp = D / "prioritaB_trovati.json"
    try:
        vecchi = json.loads(tp.read_text("utf-8"))
    except (OSError, ValueError):
        vecchi = []
    limite = (date.today() - timedelta(days=4)).isoformat()
    per = {(i["link"], tuple(i["classi"])): i for i in vecchi if (i.get("_visto") or "") >= limite}
    ora = datetime.now().isoformat(timespec="seconds")
    for it in trovati:
        it["_visto"] = ora
        per[(it["link"], tuple(it["classi"]))] = it
    tp.write_text(json.dumps(list(per.values()), ensure_ascii=False, indent=0) + "\n", "utf-8")
    S.log(f"B015/B017: avvisi trovati {len(trovati)} · in attesa {len(per)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
