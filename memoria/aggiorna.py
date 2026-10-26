#!/usr/bin/env python3
"""Aggiorna la memoria: stato per scuola (memoria/stato/scuole.json) e lista di priorità di lettura (data/priorita_lettura.json).
Unisce: data/memoria_scuole.json (stato verde/arancione/rosso, errori), data/interpelli.json (frequenza degli avvisi),
data/cms_report.json (feed RSS / ETag), data/ip_lenti.json (server lenti), data/diagnosi_errori.json (tentativi)."""
from __future__ import annotations
import json, sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from memoria import modello as M  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
J = lambda p, d: json.loads(p.read_text("utf-8")) if p.exists() else d


def main() -> int:
    oggi = date.today().toordinal()
    mem = J(ROOT / "data/memoria_scuole.json", {}).get("scuole", {})
    arch = J(ROOT / "data/interpelli.json", {}).get("interpelli", [])
    cms = J(ROOT / "data/cms_report.json", {})
    lenti = set(J(ROOT / "data/ip_lenti.json", {}).get("scuole", {}))
    diag = J(ROOT / "data/diagnosi_errori.json", {}).get("scuole", {})
    n_av, prima = defaultdict(int), {}
    for i in arch:
        c = i.get("codiceScuola")
        if c and i.get("pubblicato"):
            n_av[c] += 1
            prima[c] = min(prima.get(c, "9999"), i["pubblicato"][:10])
    inizio = min(prima.values(), default=date.today().isoformat())
    T = max(1, oggi - date.fromisoformat(inizio).toordinal())
    grp = {c: M.gruppo_scuola(e.get("nome", ""), e.get("regione", "")) for c, e in mem.items()}
    gv, gn = defaultdict(int), defaultdict(int)
    for c in mem:
        gv[grp[c]] += n_av.get(c, 0); gn[grp[c]] += 1
    glob = sum(gv.values()) / max(1, len(mem)) / T
    media = {k: max(glob, (gv[k] / T + 30 * glob) / (gn[k] + 30)) for k in gn}
    b_codici = set()
    try:
        from openpyxl import load_workbook
        for r in list(load_workbook(ROOT / "data/Elenco_scuole_B015_B017.xlsx", read_only=True)["Tutta Italia"].iter_rows(values_only=True))[1:]:
            if r and r[4]:
                b_codici.add(str(r[4]).strip().upper())
    except Exception:  # noqa: BLE001
        pass
    sede2ist = {}
    for ep in [ROOT / "data/sicilia_scuole.json"] + sorted((ROOT / "data/scuole").glob("*_elenco.json")):
        if "priorita" in ep.name:
            continue
        for c, v in J(ep, {}).get("scuole", {}).items():
            for s in v.get("sedi", []):
                sede2ist[s] = c
    domanda = defaultdict(lambda: 1.0)
    for s in b_codici:
        domanda[sede2ist.get(s, s)] = 3.0  # istituti dove B015/B017 sono plausibili: valore triplo (tue classi)
    stato = {}
    for c, e in mem.items():
        if e["stato"] == "rosso":
            continue
        d = diag.get(c, {})
        tent = len(d.get("t", [])) if d else 0
        ok = 1 if e["stato"] == "verde" else 0
        ult = e.get("ultima_lettura")
        stato[c] = {"avvisi": n_av.get(c, 0), "giorni_osservati": T, "media_gruppo": media[grp[c]], "successi": ok * max(1, tent), "tentativi": max(1, tent),
                    "ultima_lettura_giorni": date.fromisoformat(ult).toordinal() if ult else None, "feed": bool(cms.get(c, {}).get("rss")), "ip_lento": c in lenti}
    ordine = M.ordina(stato, oggi, domanda)
    (ROOT / "memoria/stato/scuole.json").write_text(json.dumps({"aggiornato": datetime.now().isoformat(timespec="seconds"), "scuole": stato}, ensure_ascii=False, separators=(",", ":")) + "\n", "utf-8")
    (ROOT / "data/priorita_lettura.json").write_text(json.dumps({"aggiornato": datetime.now().isoformat(timespec="seconds"), "ordine": [[c, round(p, 5)] for c, p in ordine[:2500]]}, separators=(",", ":")) + "\n", "utf-8")
    print(f"memoria: {len(stato)} scuole in priorita', prime: {[c for c,_ in ordine[:5]]}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
